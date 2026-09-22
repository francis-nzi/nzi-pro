"""Portal vehicle/travel workbook creation and validated row parsing."""
from __future__ import annotations
import io
import math
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from fastapi import HTTPException
from services.portal_upload_sites import row_site
from services.download_filenames import build_download_filename, safe_filename_part

LABELS = {"company_vehicles": "Company Vehicles", "business_travel": "Business Travel"}
HEADERS = ["Identifier", "Activity", "Annual Quantity", "Site Name", "Notes"] + [f"Month {i}" for i in range(1, 13)]


def activity_options(factors):
    return {f"{f['scope']} | {f['original_id']} | {f.get('report_label') or f.get('category')} | {f.get('uom') or ''}": f
            for f in factors if f.get("scope") and f.get("original_id")}


def build_workbook(meta, bucket, sites, selected, factors):
    label = LABELS[bucket]
    wb = Workbook()
    ws = wb.active
    ws.title = label
    ws.append(["Client Name:", meta.get("client_name"), None, None, "Job Number:", meta.get("job_number")])
    ws.append(["Site Name:", selected["site_name"] if selected else "Choose a site per row", None, None,
               "Reporting Period:", f"{meta.get('reporting_period_start') or '?'} to {meta.get('reporting_period_end') or '?'}"])
    ws.append(["Reporting Year:", meta.get("reporting_year")])
    ws.append(["Choose an activity (including its unit). Enter annual quantity OR monthly quantities; monthly values are summed. Use a vehicle registration or journey reference as identifier."])
    ws.merge_cells("A4:Q4")
    ws["A4"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[4].height = 32
    ws.append(HEADERS)
    ws.freeze_panes = "C6"
    for cell in ws[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for key in ("A1", "E1", "A2", "E2", "A3"):
        ws[key].font = Font(bold=True)
    for column, width in {"A": 25, "B": 65, "C": 22, "D": 35, "E": 36, "F": 28}.items():
        ws.column_dimensions[column].width = width
    for name, values, column, defined in [
        ("Sites", [s["site_name"] for s in sites], "D", "PortalSites"),
        ("Activities", list(activity_options(factors)), "B", "PortalActivities"),
    ]:
        lookup = wb.create_sheet(name)
        lookup.append([name])
        for value in values: lookup.append([value])
        lookup.column_dimensions["A"].width = 110 if name == "Activities" else 44
        if values:
            wb.defined_names.add(DefinedName(defined, attr_text=f"'{name}'!$A$2:$A${len(values)+1}"))
            validation = DataValidation(type="list", formula1=defined, allow_blank=True)
            validation.showErrorMessage = True
            validation.error = f"Choose a value from the {name} sheet."
            ws.add_data_validation(validation)
            validation.add(f"{column}6:{column}10005")
    for row in range(6, 106):
        ws.cell(row, 1).number_format = "@"
        if selected: ws.cell(row, 4, selected["site_name"])
        for col in [3, *range(6, 18)]: ws.cell(row, col).number_format = "#,##0.00"
    output = io.BytesIO()
    wb.save(output)
    filename = build_download_filename(job_number=meta.get("job_number"), client_name=meta.get("client_name"),
        descriptor=label, period_start=meta.get("reporting_period_start"), period_end=meta.get("reporting_period_end"),
        reporting_year=meta.get("reporting_year"), suffix=" " + safe_filename_part(selected["site_name"]) if selected else "")
    return output.getvalue(), filename


def parse_workbook(raw, bucket, sites, default_id, factors):
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=False)
    except Exception as exc:
        raise HTTPException(400, "Could not read workbook. Upload the completed XLSX template.") from exc
    label = LABELS[bucket]
    if label not in wb.sheetnames:
        raise HTTPException(400, f"Upload a {label} template.")
    ws = wb[label]
    if [ws.cell(5, i).value for i in range(1, 18)] != HEADERS:
        raise HTTPException(400, "Template column headings have changed. Download a fresh template.")
    options = activity_options(factors)
    rows, errors = [], []
    def number(value):
        if value is None or value == "": return None
        try: result = float(value)
        except (ValueError, TypeError): raise ValueError("Quantities must be numbers, not formulas or text")
        if not math.isfinite(result) or result < 0: raise ValueError("Quantities must be finite, non-negative numbers")
        return result
    for index in range(6, ws.max_row + 1):
        values = [ws.cell(index, col).value for col in range(1, 18)]
        if not any(v not in (None, "") for i, v in enumerate(values) if i != 3): continue
        try:
            factor = options.get(str(values[1] or "").strip())
            if factor is None: raise ValueError("Choose a valid activity from the Activities sheet")
            site_id = row_site(sites, values[3], default_id)
            months = [number(v) for v in values[5:]]
            annual = number(values[2])
            quantity = sum(v or 0 for v in months) if any(v is not None for v in months) else annual
            if quantity is None or quantity <= 0: raise ValueError("Enter an annual quantity or monthly quantities greater than zero")
            if annual is not None and any(v is not None for v in months) and not math.isclose(annual, quantity, abs_tol=0.01):
                raise ValueError("Annual quantity does not match the monthly total")
            rows.append({"scope": factor["scope"], "original_id": factor["original_id"],
                "category": factor["category"], "report_label": factor.get("report_label"),
                "uom": factor.get("uom"), "site_id": site_id, "qty": quantity,
                "identifier": str(values[0] or "").strip(), "notes": str(values[4] or "").strip(),
                **{f"month_{i+1}": value for i, value in enumerate(months)}})
        except (ValueError, HTTPException) as exc:
            errors.append({"row": index, "reason": str(exc.detail) if isinstance(exc, HTTPException) else str(exc)})
    return {"rows": rows, "errors": errors, "ready_count": len(rows)}
