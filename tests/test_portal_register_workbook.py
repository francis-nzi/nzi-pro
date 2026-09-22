from pathlib import Path
import sys
import io
from datetime import date
import pytest
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from services.portal_register_workbook import build_workbook, parse_workbook, activity_options, LABELS
from services.portal_upload_sites import row_site, selected_site
from fastapi import HTTPException

META = {"job_number": "J000663", "client_name": "Test Client", "reporting_period_start": date(2025,4,1), "reporting_period_end": date(2026,3,31), "reporting_year": 2026}
SITES = [{"site_id": 1, "site_name": "Office"}, {"site_id": 2, "site_name": "Factory"}]
FACTORS = [{"scope": "Scope 1", "original_id": "V1", "report_label": "Diesel car", "category": "Company Vehicles", "uom": "km"}]


def workbook(bucket):
    raw, filename = build_workbook(META, bucket, SITES, SITES[0], FACTORS)
    return load_workbook(io.BytesIO(raw)), filename


def save(wb):
    stream = io.BytesIO(); wb.save(stream); return stream.getvalue()


@pytest.mark.parametrize("bucket", LABELS)
def test_template_filename_metadata_prefill_and_import(bucket):
    wb, filename = workbook(bucket)
    assert filename == f"J000663 Test Client {LABELS[bucket]} 2025-2026 Office.xlsx"
    ws = wb[LABELS[bucket]]
    assert ws["B1"].value == "Test Client"
    assert ws["F1"].value == "J000663"
    assert ws["B2"].value == "Office"
    assert ws["F2"].value == "2025-04-01 to 2026-03-31"
    assert all(ws.cell(row,4).value == "Office" for row in range(6,106))
    assert len(ws.data_validations.dataValidation) == 2
    assert parse_workbook(save(wb),bucket,SITES,None,FACTORS)["ready_count"] == 0
    ws["A6"] = "AB12 CDE"; ws["B6"] = next(iter(activity_options(FACTORS))); ws["C6"] = 500
    ws["A7"] = "Trip 2"; ws["B7"] = ws["B6"].value; ws["D7"] = "Factory"; ws["F7"] = 20; ws["G7"] = 30
    result = parse_workbook(save(wb),bucket,SITES,None,FACTORS)
    assert not result["errors"]
    assert [(row["site_id"],row["qty"]) for row in result["rows"]] == [(1,500),(2,50)]
    assert result["rows"][1]["month_1"] == 20


@pytest.mark.parametrize("field,value,reason", [("D6","Foreign site","permitted"),("B6","Invalid activity","activity"),("C6",-1,"non-negative"),("C6","=1+2","numbers")])
def test_invalid_rows_are_rejected(field,value,reason):
    wb,_ = workbook("company_vehicles"); ws=wb["Company Vehicles"]
    ws["B6"] = next(iter(activity_options(FACTORS))); ws["C6"] = 20
    ws[field] = value
    result = parse_workbook(save(wb),"company_vehicles",SITES,None,FACTORS)
    assert result["ready_count"] == 0
    assert reason in result["errors"][0]["reason"]


def test_site_permissions_and_missing_site():
    for callback in [lambda: selected_site(SITES,999), lambda: row_site(SITES,"",None), lambda: row_site(SITES,"Other",1)]:
        with pytest.raises(HTTPException): callback()
    assert row_site(SITES,"",2) == 2


def test_register_import_uses_existing_register_and_review(monkeypatch):
    import api.portal_data_entry_routes as routes
    class Conn:
        def execute(self,sql,params=None):
            self.sql=sql; self.params=params
            if "INSERT INTO job_emission_sources" in sql: self.insert=(sql,params)
            return self
        def fetchone(self): return (42,)
    for name in ("_ensure_job_scope_rows_schema","ensure_portal_data_entry_schema","_ensure_emission_register_schema","_assert_data_entry_open","record_audit_event"):
        monkeypatch.setattr(routes,name,lambda *args,**kwargs: None)
    monkeypatch.setattr(routes,"_resolve_job_or_404",lambda *args: 663)
    monkeypatch.setattr(routes,"load_bucket_category_map",lambda *args: {})
    monkeypatch.setattr(routes,"bucket_for_category",lambda mapping,category: category)
    monkeypatch.setattr(routes,"_resolve_scope_row_factor_for_creation",lambda *args: (1,2,0.1,"tCO2e"))
    monkeypatch.setattr(routes,"_lookup_factor_from_reference",lambda *args: {})
    for bucket,source in [("company_vehicles","asset"),("business_travel","business_travel")]:
        con=Conn()
        result=routes._create_portal_data_entry_row(con,None,bucket,{"category":bucket,"scope":"Scope 1","original_id":"V1","site_id":1,"qty":10,"uom":"km"},{"client_db_id":5})
        sql,params=con.insert
        assert params[3] == source
        assert params[4] == 1
        assert "FALSE, 'pending_review', TRUE" in sql
        assert result["review_status"] == "pending_review"


def test_commuting_site_prefill_and_parse():
    import api.employee_commuting_routes as commuting
    wb=load_workbook(io.BytesIO(commuting._build_template_workbook(META,"Office",sites=SITES)))
    ws=wb[commuting.COMMUTING_SHEET]; wfh=wb[commuting.WFH_SHEET]
    assert ws["B6"].value == wfh["B6"].value == "Office"
    assert ws.cell(commuting.DATA_START_ROW,10).value == "Office"
    assert wfh.cell(commuting.DATA_START_ROW,6).value == "Office"
    assert commuting._parse_template(save(wb)) == []
    ws.cell(commuting.DATA_START_ROW,1,"EMP-1")
    ws.cell(commuting.DATA_START_ROW,2,"Car")
    ws.cell(commuting.DATA_START_ROW,10,"Factory")
    rows=commuting._parse_template(save(wb))
    assert rows[0]["site_name"] == "Factory"


def test_commuting_groups_by_validated_row_sites(monkeypatch):
    import api.portal_commuting_routes as routes
    monkeypatch.setattr(routes,"_parse_template",lambda raw: [
        {"employee_name":"EMP-1","site_name":"Office"}, {"employee_name":"EMP-2","site_name":"Factory"}])
    calls=[]
    def resolve(con,job,site,entries):
        calls.append(site)
        return {"ready_rows":[{**entry,"site_id":site,"calc_tco2e":0.1} for entry in entries],"unresolved_rows":[]}
    monkeypatch.setattr(routes,"_resolve_manual_commuting_rows",resolve)
    monkeypatch.setattr(routes,"_assert_employee_id_unique",lambda *args: None)
    result=routes._portal_upload_preview(None,663,None,b"",sites=SITES)
    assert calls == [1,2]
    assert result["ready_count"] == 2


def test_register_upload_commits_atomically_and_blocks_invalid_rows(monkeypatch):
    import asyncio
    import api.portal_data_entry_routes as routes
    class Conn:
        def __enter__(self): return self
        def __exit__(self,kind,*args): exits.append(kind)
    exits=[]; modes=[]; saved=[]
    def connect(**kwargs): modes.append(kwargs); return Conn()
    monkeypatch.setattr(routes,"get_conn",connect)
    monkeypatch.setattr(routes,"_register_workbook_context",lambda *args: (663,META,SITES,SITES[0],FACTORS))
    monkeypatch.setattr(routes,"_assert_data_entry_open",lambda *args: None)
    wb,_=workbook("company_vehicles"); ws=wb["Company Vehicles"]
    ws["B6"]=next(iter(activity_options(FACTORS))); ws["C6"]=100
    async def read(file): return save(wb)
    monkeypatch.setattr(routes,"_read_register_upload",read)
    monkeypatch.setattr(routes,"_create_portal_data_entry_row",lambda con,request,bucket,row,user: saved.append(row))
    result=asyncio.run(routes.portal_register_upload_commit(None,"company_vehicles",None,1,{"client_db_id":5}))
    assert result["inserted"] == 1
    assert modes == [{"autocommit":False}]
    assert exits == [None]
    saved.clear(); ws["B7"]="Unknown activity"; ws["C7"]=100
    with pytest.raises(HTTPException):
        asyncio.run(routes.portal_register_upload_commit(None,"company_vehicles",None,1,{"client_db_id":5}))
    assert saved == []
    assert exits[-1] is HTTPException


def test_commuting_download_uses_selected_permitted_site(monkeypatch):
    import api.portal_commuting_routes as routes
    class Conn:
        def __enter__(self): return self
        def __exit__(self,*args): pass
    monkeypatch.setattr(routes,"get_conn",lambda **kwargs: Conn())
    monkeypatch.setattr(routes,"_resolve_job_or_404",lambda *args: 663)
    monkeypatch.setattr(routes,"_job_meta",lambda *args: META)
    monkeypatch.setattr(routes,"permitted_sites",lambda *args: SITES)
    response=routes.portal_commuting_template({"client_db_id":5},site_id=2)
    assert response.headers["x-filename"] == "J000663 Test Client Employee Commuting 2025-2026 Factory.xlsx"
    wb=load_workbook(io.BytesIO(response.body))
    assert wb["Employee Commuting"]["B6"].value == "Factory"
    with pytest.raises(HTTPException): routes.portal_commuting_template({"client_db_id":5},site_id=999)
